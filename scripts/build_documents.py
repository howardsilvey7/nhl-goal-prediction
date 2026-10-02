"""Build three formal DOCX documents from frozen results and authored text.

PDFs are exported from these same DOCX sources by export_documents.ps1.
"""
from pathlib import Path
import sys
import json
import re
import hashlib
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs'
OUT.mkdir(exist_ok=True)
sys.path.insert(0,str(ROOT/'src'))
from nhl_forecast.data import checksum
M=pd.read_csv(ROOT/'reports/holdout-metrics.csv').set_index('model')
D=pd.read_csv(ROOT/'reports/development-metrics.csv')
S=pd.read_csv(ROOT/'reports/data-summary.csv')
C=pd.read_csv(ROOT/'reports/holdout-paired-comparisons.csv').set_index('model')
F=json.loads((ROOT/'reports/frozen-models.json').read_text())
NAMES={'league_poisson':'League Poisson','rolling_poisson':'Rolling Poisson',
 'ridge_poisson_proxy':'Ridge Poisson proxy','poisson':'Independent Poisson',
 'negative_binomial':'Negative binomial','attack_defence_poisson':'Attack and defence Poisson',
 'bivariate_poisson':'Bivariate Poisson', 'boosted_poisson':'Poisson boosting',
 'random_forest_poisson_proxy':'Forest Poisson proxy'}
PRIMARY=F['selected_primary']; P=M.loc[PRIMARY]; BASE=M.loc['league_poisson']; PC=C.loc[PRIMARY]
REFERENCES=[
 '[1] Silvey, H. (2020). National Hockey League Data: Analysis and Prediction. Undergraduate honours project, Carleton University, School of Mathematics and Statistics. Supervisor: Dr. Jason Nielsen. April 7, 2020. https://carleton.ca/math/wp-content/uploads/Howard-Silvey-Honours-Project-2.pdf',
 '[2] Gneiting, T., and Raftery, A. E. (2007). Strictly Proper Scoring Rules, Prediction, and Estimation. Journal of the American Statistical Association, 102(477), 359-378. https://doi.org/10.1198/016214506000001437',
 '[3] Karlis, D., and Ntzoufras, I. (2003). Analysis of sports data by using bivariate Poisson models. Journal of the Royal Statistical Society, Series D, 52(3), 381-393. https://doi.org/10.1111/1467-9884.00366',
 '[4] Dixon, M. J., and Coles, S. G. (1997). Modelling Association Football Scores and Inefficiencies in the Football Betting Market. Journal of the Royal Statistical Society, Series C, 46(2), 265-280. https://doi.org/10.1111/1467-9876.00065',
 '[5] Macdonald, B., Lennon, C., and Sturdivant, R. (2012). Evaluating NHL Goalies, Skaters, and Teams Using Weighted Shots. arXiv:1205.1746. https://arxiv.org/abs/1205.1746',
 '[6] Thomas, A. C., Ventura, S. L., Jensen, S., and Ma, S. (2013, revised preprint). Competing Process Hazard Function Models for Player Ratings in Ice Hockey. arXiv:1208.0799v2. https://arxiv.org/abs/1208.0799',
 '[7] Sadeghkhani, A., and Ahmed, S. E. (2019). Predicting the scoring time in hockey. arXiv:1903.10889. https://arxiv.org/abs/1903.10889',
 '[8] MoneyPuck.com (2026). Download Data. All-team game-by-game data, snapshot retrieved October 2, 2026. Non-commercial research use with attribution. https://moneypuck.com/data.htm',
 '[9] National Hockey League (2026). GameCenter landing responses. Official period scoring used for 54 stratified verification cases. Endpoint pattern: https://api-web.nhle.com/v1/gamecenter/{game_id}/landing',
 '[10] scikit-learn developers (2026). PoissonRegressor and HistGradientBoostingRegressor documentation, version 1.9.1. https://scikit-learn.org/stable/modules/generated/sklearn.linear_model.PoissonRegressor.html ; https://scikit-learn.org/stable/modules/generated/sklearn.ensemble.HistGradientBoostingRegressor.html',
 '[11] Silvey, H. R. (2026). NHL Honours Project 2020 Audit and 2026 Research Baseline. Version 2.0, September 27, 2026. Preserved transfer document in docs/TRANSFER_BASELINE.md.',
 '[12] Silvey, H. R. (2026). NHL Goal Prediction research repository. Research edition 0.1.0. https://github.com/howardsilvey7/nhl-goal-prediction']

def new_doc(title,subtitle):
 doc=Document();sec=doc.sections[0]
 sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=sec.bottom_margin=Inches(.85)
 sec.left_margin=sec.right_margin=Inches(.9)
 normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
 normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.15
 for name,size in [('Title',22),('Subtitle',13),('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
  style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
  style.paragraph_format.space_before=Pt(12);style.paragraph_format.space_after=Pt(8)
  if name.startswith('Heading'):style.paragraph_format.keep_with_next=True
 # The bundled default template has theme fonts and a Title border. Clear both.
 for style in doc.styles:
  for border in list(style._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
  for fonts in style._element.iter(qn('w:rFonts')):
   for attr in list(fonts.attrib):
    if 'theme' in attr.lower():del fonts.attrib[attr]
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
  fonts=doc.styles[name]._element.get_or_add_rPr().get_or_add_rFonts()
  for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),'Times New Roman')
 doc.core_properties.author='Howard R. Silvey';doc.core_properties.title=title
 doc.core_properties.subject='2026 revision and extension of the 2020 NHL honours project'
 doc.core_properties.keywords='NHL, regulation goals, temporal validation, reproducible research'
 doc.add_paragraph(title,'Title');doc.add_paragraph(subtitle,'Subtitle')
 doc.add_paragraph('Howard R. Silvey\nResearch edition 0.1.0\nOctober 2, 2026')
 doc.add_paragraph('Original honours project supervised by Dr. Jason Nielsen\nCarleton University School of Mathematics and Statistics\nOriginal submission April 7, 2020')
 doc.add_paragraph('Independent 2026 research revision and extension. This edition is separate from the original university submission and does not imply institutional review or supervisor endorsement of the new analysis.')
 doc.add_paragraph('Primary sharing format PDF\nEditable companion DOCX\nRepository github.com/howardsilvey7/nhl-goal-prediction')
 footer=sec.footer.paragraphs[0];footer.alignment=2
 r=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
 footer.style=doc.styles['Normal'];footer.paragraph_format.space_after=Pt(0)
 doc.add_page_break()
 return doc

def p(doc,text):
 paragraph=doc.add_paragraph(text)
 paragraph.paragraph_format.keep_together=True

def equation(doc,text):
 paragraph=doc.add_paragraph();paragraph.alignment=1
 paragraph.paragraph_format.keep_together=True
 for part in re.split(r'(\{[^}]+\})',text):
  if not part:continue
  run=paragraph.add_run(part[1:-1] if part.startswith('{') else part)
  run.font.name='Times New Roman';run.font.size=Pt(12);run.italic=True
  if part.startswith('{'):run.font.subscript=True

def page_break(doc):
 if doc.paragraphs and doc.paragraphs[-1]._p.xpath('.//w:br[@w:type="page"]'):return
 doc.add_page_break()

def section(doc,title,paragraphs,level=1,page=False):
 if page:page_break(doc)
 doc.add_heading(title,level)
 for text in paragraphs:p(doc,text)

def table(doc,headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for cell,head in zip(t.rows[0].cells,headers):cell.text=str(head)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for values in rows:
  cells=t.add_row().cells
  for c,v in zip(cells,values):c.text=str(v)
 for row in t.rows:
  for cell in row.cells:
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(4);para.paragraph_format.line_spacing=1
    para.paragraph_format.keep_with_next=True
    for run in para.runs:
     run.font.size=Pt(9);run.font.name='Times New Roman';run.font.color.rgb=RGBColor(0,0,0)
  cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
 if widths:
  t.autofit=False
  for row in t.rows:
   for cell,w in zip(row.cells,widths):cell.width=Inches(w)
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  node=OxmlElement('w:'+edge);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'CBCBCB');borders.append(node)
 t._tbl.tblPr.append(borders)
 # Neutral grayscale table design rather than default Office blue.
 for cell in t.rows[0].cells:
  tcPr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E7E7');tcPr.append(shade)
  for para in cell.paragraphs:
   for run in para.runs:run.bold=True;run.font.color.rgb=RGBColor(0,0,0)
 doc.add_paragraph()
 return t

def figure(doc,name,caption):
 doc.add_picture(str(ROOT/'reports/figures'/name),width=Inches(6.6))
 doc.paragraphs[-1].paragraph_format.keep_with_next=True
 para=doc.add_paragraph(caption);para.paragraph_format.space_after=Pt(10)
 for r in para.runs:r.font.size=Pt(9);r.italic=True

def contents(doc,entries):
 doc.add_heading('Contents',1)
 for line in entries:p(doc,line)
 p(doc,'The PDF includes bookmarks for the chapter headings. Table and figure captions identify the numbered research outputs.')
 page_break(doc)

def save(doc,name):
 path=OUT/f'{name}.docx';doc.save(path);print(path)

def thesis():
 doc=new_doc('National Hockey League Goal Prediction','Modern Statistical and Machine Learning Approaches\nA Revision and Extension of the 2020 Undergraduate Honours Project')
 section(doc,'Author preface',[
  'This revision returns to the paired-score forecasting idea proposed in my 2020 honours project. I preserve the original submission as an archival work and examine its limitations before introducing a separate game-level experiment. The aim is to improve the research design, report uncertainty honestly and make the computational analysis reproducible.',
  'The 2020 project arose within an overloaded final-semester academic schedule. That context helps explain its limited scope but does not change the standards used to assess its statistical claims. The revision recognizes the original project’s useful diagnosis of its own aggregation problem and its anticipation of joint home and away goal modeling.',
  'The present edition reports a completed first empirical comparison, not completion of every item in the broader research programme. Original source-code recovery, independent academic review, more extensive availability verification and advanced extensions remain open. All new numerical results belong to 2026, not to the 2020 submission.',
  'Computational assistance from OpenAI Codex supported source inspection, software implementation, numerical checks and document preparation. Authorship remains Howard R. Silvey. Automated verification and software agreement do not substitute for independent scientific review.'])
 section(doc,'Abstract',[
  f'The original NHL honours project used 123 team-season rows to address an individual-game scoring question. This revision audits that mismatch and conducts a separate retrospective forecasting study of {int(S.n_games.sum()):,} NHL regular-season games from 2016-17 through 2024-25. Paired regulation goals are reconstructed from MoneyPuck’s actual-goal counts and game duration; 54 stratified cases agree with official period scoring. One incomplete source record is excluded.',
  'Nine baselines and model variants are compared using expanding temporal windows. Hyperparameter choices use the most recent training season as an inner validation period. Five outer development seasons determine the model family; the 2024-25 season is evaluated after configurations are frozen. Feature histories include only strictly earlier source calendar dates. Exact historical publication-time availability remains unverified.',
  f'The frozen bivariate-Poisson family is numerically tied with independent Poisson in development and estimates a zero shared component. On {int(P["n"]):,} holdout games, its joint negative log likelihood is {P.joint_nll:.4f}, compared with {BASE.joint_nll:.4f} for the expanding league baseline. The paired calendar-week bootstrap difference is {PC.mean_difference:.4f}, with a 95% interval [{PC.low95:.4f}, {PC.high95:.4f}]. Combined mean absolute error changes from {BASE.combined_mae:.4f} to {P.combined_mae:.4f} goals per team-game. A ridge model with Poisson margins has a slightly smaller holdout log score but does not replace the frozen family.',
  'The evidence supports modest improvements in probabilistic forecasts from regularized recent-form predictors. It does not support a positive-dependence advantage or a large point-forecast gain. Original historical results are verified as reported, not computationally reproduced. The repository, editable manuscript and primary PDF preserve these distinctions.'])
 contents(doc,['1 Introduction','2 The 2020 honours project and archival audit','3 Related research and statistical framework',
  '4 Data and prediction contract','5 Temporal validation and methods','6 Development results',
  '7 Frozen final holdout results','8 Calibration uncertainty and reproducibility',
  '9 Discussion and limitations','10 Conclusion and remaining programme','References','Appendices'])
 section(doc,'1 Introduction',[
  'Forecasting hockey scores requires a distinction between describing a completed season and predicting a particular future game. A season scoring average summarizes a team’s realized performance; it does not directly identify the distribution of goals against an opponent on a given date. The two tasks have different units, information sets and validation requirements.',
  'A useful pregame forecast should express the probabilities of possible home and away scores. Expected goals alone do not answer how likely a regulation draw is, how uncertain the total is or how much probability should be assigned to an unusual scoreline. A paired predictive distribution supports these quantities while making forecast uncertainty explicit.',
  'The broad research question is which statistical or machine-learning framework most accurately and reliably estimates the joint distribution of future home and away goals using only information available before a game. The first experiment operationalizes that question with earlier calendar-date histories, a fixed set of practical models and a clearly separated final season.',
  'The contributions are an archival methodological audit, a canonical paired-game pipeline, tests for temporal leakage, nested season-based model selection, proper probabilistic evaluation, a frozen holdout comparison and an independent base-R check of the Python Poisson fit. These contributions are complementary: none alone establishes a deployable forecasting system.',
  'The study is predictive and retrospective. It does not estimate causal effects of shots, rest or team strength, and it does not study betting returns. Its scores assess the stated target and validation design, not general superiority across every NHL era or information source.'],page=True)
 section(doc,'2 The 2020 honours project and archival audit',[
  'The original project was submitted on April 7, 2020 under the title National Hockey League Data: Analysis and Prediction [1]. It considered ordinary multiple regression, forward selection, backward elimination, residual plots and leave-one-out cross-validation. Its stated interest was home-team game scoring and a later extension to a vector-valued home/away response.',
  'Its effective estimation sample was 123 team-season records: 30 teams for 2016-17 and 31 for each of the following three seasons. Approximately 40 candidate columns produced roughly 4,920 table cells. The cells are not independent regression observations. The source’s cross-validation screenshot explicitly reports 123 samples and five predictors.',
  'The model response was a season-level goals quantity, converted to a nominal per-game estimate by dividing a predicted total by 82. That division does not create opponent, venue or game-state information. The paper itself recognized clustered forecasts and recommended game-level modeling.',
  'Several same-season candidate measures were contemporaneous with, or partly constructed from, the goal response. Shooting percentages, adjusted goals, high-danger goals, power-play performance, standings and save measures can be informative when lagged; when calculated over the target period they do not represent a historical pregame information set.',
  'The original PDF survives locally and is preserved byte-for-byte in the archive. A targeted filename inventory of local University, Documents and Downloads materials found no original NHL R scripts or underlying team-season table. This search does not establish that every device, cloud location or archive lacks them. Exact reproduction remains unresolved; a new game-level study is not a substitute for the missing original inputs.'],page=True)
 table(doc,['Historical quantity','Reported value','Evidence and status'],[
  ('Regression sample','123 team-seasons','PDF physical pages 17 and 39; source-verified'),
  ('Backward model R squared','0.1161','PDF physical pages 28 and 32; source-verified'),
  ('Forward model R squared','0.06451','PDF physical page 32; source-verified'),
  ('LOOCV RMSE','46.20022','PDF physical page 39 screenshot; source-verified'),
  ('LOOCV R squared','0.03927991','PDF physical page 39 screenshot; source-verified'),
  ('LOOCV MAE','37.91543','PDF physical page 39 screenshot; source-verified')],[1.8,1.3,3.5])
 p(doc,'Table 1. Historical numbers remain on the original team-season response scale. They are transcribed and visually checked, not independently recomputed. They are not directly comparable with modern regulation-goal errors.')
 section(doc,'2.1 Corrections and preserved insights',[
  'LOOCV across pooled team-seasons estimates interpolation under assumptions unlike future-season deployment. Later seasons can influence earlier ones, repeated teams induce dependence and the original record does not establish that selection was repeated inside each validation fit. Stepwise significance does not provide a reliable guarantee of future forecast quality.',
  'The ordinary least-squares assumptions require precise treatment. The Central Limit Theorem does not establish independent, homoscedastic or normally distributed errors. A residual Q-Q plot addresses a different issue from temporal dependence. Adjusted R squared penalizes model size but does not remove all optimism from searching many candidate models.',
  'The original’s strongest contribution is its honest recognition that the chosen linear model performed poorly and that game-level paired responses, regularization and richer information could improve the design. This revision treats those ideas as hypotheses to evaluate rather than outcomes that must be confirmed.'],level=2,page=True)
 section(doc,'3 Related research and statistical framework',[
  'The theoretical evaluation framework follows proper scoring rules [2]. A logarithmic score rewards a forecast distribution for assigning probability to the realized outcome. A mean negative log likelihood is therefore an appropriate primary loss for paired discrete scores. It addresses a different question from minimizing absolute point error.',
  'Bivariate Poisson models introduce a shared latent count to represent positive dependence [3]. The construction is useful as a testable score-distribution model, but a shared Poisson component cannot represent negative conditional covariance. Marginal scoring environments and conditional game interactions must be distinguished before interpreting an unconditional score correlation.',
  'Dixon and Coles developed a football-specific Poisson extension addressing dynamic performance and low-score behavior [4]. It provides methodological context, not direct evidence that its correction applies to hockey. Higher scoring, empty-net strategies, penalties and overtime rules motivate separate NHL diagnostics before adopting such a correction.',
  'Macdonald, Lennon and Sturdivant use shot-level logistic estimates to evaluate NHL performance [5]. Their work supports examining shot-quality measures as complements to traditional quantities. It does not establish that retrospective expected-goals files were available at every historical forecast date; that availability issue is why xG is outside this primary experiment.',
  'Thomas and colleagues model competing scoring processes with player effects and shrinkage [6]. Sadeghkhani and Ahmed examine predictive densities for scoring times [7]. These studies show that event timing and on-ice context can support richer models, while differing from the present pregame paired-score task. Neither model is implemented here.',
  'The implemented model ladder begins with league and recent-form baselines, then regularized count and point models, a categorical attack/defence model, a plug-in shared-count extension and controlled tree ensembles. This is a selected research comparison rather than an exhaustive survey. The related-work review is focused and does not claim systematic coverage of all publications through 2026.'],page=True)
 section(doc,'3.1 Distributional definitions',[
  'For game g, the target is the pair of home and away regulation goal counts. The predictive distribution conditions on the history admitted by the information boundary. Its conditional means can differ by home role, team form and opponent form.',
  'Independent Poisson factorizes the paired distribution into two Poisson margins. Each fitted count-model mean uses a log link. The NB2 alternative allows conditional variance to exceed the mean by a nonnegative quadratic dispersion term estimated from training data.',
  'The plug-in bivariate construction adds the same latent Poisson count to two independent side-specific counts. Their means are the home mean minus c, the away mean minus c and c. The shared count preserves both marginal means and gives conditional covariance c. Rho is estimated on training data with marginal predictions held fixed; this is not full joint maximum-likelihood regression.',
  'Ridge and forest predictions are clipped below at 0.05 goals and embedded in independent Poisson margins for probability comparison. Their mean estimators optimize squared-error objectives, not count likelihood. Their probability outputs are explicit Poisson proxies and should not be described as learned full distributions.'],level=2,page=True)
 equation(doc,'Y{g} = (Y{gH}, Y{gA})')
 equation(doc,'p{g}(h, a | I{g}) = Pois(h; μ{gH}) Pois(a; μ{gA})')
 equation(doc,'Var(Y | I) = μ + α μ²')
 equation(doc,'Y{H} = U + W,     Y{A} = V + W')
 equation(doc,'c = ρ min(μ{H}, μ{A}),     0 ≤ ρ ≤ 0.8')
 section(doc,'4 Data and prediction contract',[
  'The primary modern source is MoneyPuck’s publicly listed all-team game-by-game file [8]. The retrieved snapshot is immutable locally and identified by SHA256. Only original code, aggregate research outputs and provenance information are published. Raw files and row-level derived datasets are excluded from the public repository.',
  'The study selects regular-season records for season-start years 2016 through 2024, with situation equal to all. A canonical game pairs one home and one away record. Paired date, season, opponents, elapsed time and mirrored goals must agree. Duplicate roles, missing pairs and inconsistent mirrored values cause a failure rather than an unreported repair.',
  'MoneyPuck goals are actual played goals rather than the shootout scoreboard marker. Regulation reconstruction uses all-situation elapsed time. A 3,600-second game retains its played score; a longer game subtracts the actual sudden-death goal from the winning team. A tied played score after overtime requires no subtraction and is treated as a shootout game. The resulting regulation score is tied whenever overtime occurred.',
  'One record, game 2019020876, contains only 1,200 elapsed seconds in the source despite appearing in a regular-season table. Both team rows have that incomplete duration. The record is excluded under the duration rule. The exclusion is a source-integrity decision, not an outcome-based choice favoring a model.',
  'Fifty-four deterministic verification cases cover two games for each finish type in each included season. Reconstructed regulation counts match official period-by-period scoring in every sampled case [9]. This validates those cases and the inferred convention; it does not prove every row in the complete dataset is correct.'],page=True)
 table(doc,['Season','Games','Home mean','Away mean'],[(f'{int(r.season)}-{str(int(r.season)+1)[-2:]}',f'{int(r.n_games):,}',f'{r.home_mean:.3f}',f'{r.away_mean:.3f}') for r in S.itertuples()],[1.5,1.1,2.,2.])
 p(doc,'Table 2. Included game counts and regulation-goal means. The 2019-20 count reflects the one incomplete-source exclusion. Shortened and disrupted seasons remain explicit temporal folds.')
 figure(doc,'scoring-by-season.png','Figure 1. Scoring environments vary over time. Each season is summarized after the study configurations were frozen; holdout summaries did not guide model selection.')
 section(doc,'4.1 Features and historical availability',[
  'All features are emitted before any current-date game updates history. Every game sharing the same source calendar date belongs to one batch. The history includes strictly earlier dates, a stronger separation than a simple row shift when multiple games occur on one day.',
  'The recent-form window contains up to 20 preceding included regular-season games. Ten fixed prior games provide shrinkage at 2.7 goals, 30 shots and 55 attempts per side. Histories carry between seasons. Expansion or renamed teams use the pooled cold-start prior; the first version treats ARI and UTA as different labels rather than imposing an unverified franchise continuation.',
  'Predictors include a home indicator, team and opponent goals for/against, shots for/against, team attempts for and opponent attempts against, days since the previous included game capped at seven, history size and an expanding side-specific league scoring rate. Goals use regulation targets, while historical shot and attempt totals may include overtime. This exposure mismatch is disclosed rather than hidden.',
  'The source gameDate field supplies a calendar ordering, not an exact UTC prediction timestamp. The pipeline therefore verifies date exclusion without pretending to verify historical publication times. Current downloadable data may contain later corrections. The study should be read as a conservative retrospective date-based forecast comparison, with a further operational audit required for live pregame deployment.'],level=2)
 section(doc,'5 Temporal validation and methods',[
  'The foundation contains 2016-17 through 2018-19. Outer validation seasons are 2019-20 through 2023-24. Each outer model is fitted using earlier seasons only. Within its training period, the most recent earlier season is an inner validation window for the fixed hyperparameter search; all preceding seasons form the inner fit.',
  'Scalers, encoders and coefficients are fitted separately within each training window. The chosen inner configuration is refitted using the complete outer training set. Coefficients remain fixed during the outer season, while rolling history and league-rate predictors can update from games on earlier dates in that season. This emulates sequential forecasting with a season-level refit schedule.',
  'The final 2024-25 season is evaluated after development selects a model family and all final configurations are frozen. Final hyperparameters use 2023-24 as inner validation and training through 2022-23, followed by a refit through 2023-24. No feature or tuning choice is changed after the holdout scores are observed. The 2025-26 season remains reserved for a future extension.',
  'The protocol was committed before development comparisons. It is a retrospective repository protocol, not an externally registered study. The final family selector used the smallest game-weighted development NLL. Independent and bivariate Poisson achieved an exact numerical tie; the implementation’s ordering selected the bivariate label. Its fitted shared component is zero, so that choice is not evidence favoring dependence.'],page=True)
 figure(doc,'temporal-validation.png','Figure 2. Each evaluation season follows its training period. Inner model selection occurs entirely within the shaded training window. The final holdout was used only after freezing configurations.')
 table(doc,['Model','Estimator and frozen search'],[
 ('League Poisson','Expanding earlier-date home/away means'),('Rolling Poisson','Smoothed attack and opponent-defence average'),
 ('Ridge proxy','L2 penalties 1, 10, 100; Poisson probability margins'),('Independent Poisson','Log link; L2 penalties 0.01, 0.1, 1'),
 ('Negative binomial','NB2 log link; penalties 0.01, 0.1; estimated dispersion'),
 ('Attack and defence','Team/opponent indicators; shared home effect; L2 shrinkage'),
 ('Bivariate Poisson','Poisson means plus training-estimated nonnegative shared component'),
 ('Poisson boosting','7 or 15 leaves; 100 or 200 iterations; no random early-stopping split'),
 ('Forest proxy','250 trees; minimum leaf 20 or 50; Poisson probability margins')],[2.1,4.5])
 p(doc,'Table 3. Each model’s search space was fixed before comparison. The penalized attack/defence model supplies shrinkage but does not produce a Bayesian posterior. The tree implementations follow the documented software API [10].')
 section(doc,'5.1 Scores and comparison uncertainty',[
  'Primary loss is mean joint negative log likelihood per game, measured using natural logarithms. Lower values are better. Joint NLL uses the model’s untruncated count probability at the realized score. MAE and RMSE are reported on a per-team-game scale, with side-specific errors and bias retained in the aggregate tables.',
  'Regulation home/draw/away probabilities are derived by summing joint score mass on a 0 to 35 grid. The implementation rejects a grid with material omitted probability mass. The Brier score is the mean sum of squared errors across the three outcome probabilities; it is not divided by the number of classes.',
  'Central 90% marginal intervals use discrete 5th and 95th percentiles. Their attainable coverage need not equal 90% exactly. Coverage is reported with interval width so that broad intervals are not mistaken for better information. Interval coverage alone is not a complete calibration test.',
  'Paired bootstrap comparisons resample whole calendar weeks, preserving each game’s model pairing. Two thousand replicates form percentile intervals for NLL differences. Week blocks reduce within-week dependence but do not remove all repeated-team dependence across weeks. The intervals are conditional descriptive comparisons for this single final season, not universal performance guarantees.'],level=2)
 section(doc,'6 Development results',[
  'The five outer seasons contain 5,885 evaluation games. Development comparisons consistently show that recent-form regularized predictors can improve joint score over a pooled league-rate forecast, while gains depend on the season. This variation matters more than a single pooled ranking.',
  'The highest-scoring environment is not necessarily the most predictable. Model error reflects both changes in mean scoring and the residual variation around each forecast. Seasonal refitting and expanding league rates address only part of that variation.',
  'The bivariate extension returns zero shared dependence in every reported outer fit. It therefore has the same joint probabilities and NLL as independent Poisson. The negative-binomial dispersion is small in the final fit and does not yield a clear improvement. More distributional flexibility has not earned a distinct advantage in this implementation.'],page=True)
 weighted=D.groupby('model').apply(lambda g:(g.joint_nll*g.n).sum()/g.n.sum(),include_groups=False).sort_values()
 table(doc,['Model','Development joint NLL','Evaluation games'],[(NAMES[name],f'{v:.5f}','5,885') for name,v in weighted.items()],[3.3,2.1,1.2])
 p(doc,'Table 4. Game-weighted mean joint NLL across the five outer development seasons. The equal independent/bivariate scores represent the same forecasts rather than two distinct successful models.')
 figure(doc,'development-folds.png','Figure 3. Selected development curves retain the changing difficulty of the temporal folds. All nine variants are tabulated in the repository, including less favorable results.')
 section(doc,'7 Frozen final holdout results',[
  f'The final holdout contains {int(P["n"]):,} games from 2024-25. The frozen bivariate label corresponds to independent Poisson because rho = {P.rho:.1f}. Its joint NLL is {P.joint_nll:.5f} versus {BASE.joint_nll:.5f} for the league baseline, an absolute reduction of {BASE.joint_nll-P.joint_nll:.5f} nats per game.',
  f'The point-error difference is small. Combined MAE is {P.combined_mae:.5f} versus {BASE.combined_mae:.5f}, a reduction of {BASE.combined_mae-P.combined_mae:.5f} goals per team-game. Combined RMSE is {P.combined_rmse:.5f} versus {BASE.combined_rmse:.5f}. The appropriate interpretation is a modest probability-quality improvement, not accurate prediction of exact future scores.',
  'The ridge Poisson proxy has the smallest numerical holdout NLL, narrowly ahead of the frozen count model. That observation does not authorize changing the primary family after inspecting the final season. The difference is reported alongside the model’s probability-assumption limitation.',
  'Poisson boosting and the forest proxy improve NLL relative to the league baseline numerically, but their paired bootstrap intervals cross zero. The categorical attack/defence model provides a smaller and uncertain gain. These findings constrain claims that machine learning or a larger model automatically improves forecasts.'],page=True)
 table(doc,['Model','Joint NLL','MAE','RMSE'],[(NAMES[name],f'{row.joint_nll:.5f}',f'{row.combined_mae:.5f}',f'{row.combined_rmse:.5f}') for name,row in M.iterrows()],[3.2,1.15,1.1,1.15])
 p(doc,'Table 5. All frozen variants on the same final-season games. MAE and RMSE average the two sides on the stated team-game scale; joint NLL is per paired game.')
 figure(doc,'holdout-comparison.png','Figure 4. Joint-NLL differences against the league baseline with 95% paired calendar-week bootstrap intervals. A negative difference favors the candidate; intervals are not adjusted for multiple comparisons.')
 section(doc,'7.1 Outcome probabilities and interval coverage',[
  f'The frozen model’s regulation-outcome Brier score is {P.outcome_brier:.5f}, compared with {BASE.outcome_brier:.5f} for the league forecast. Its outcome log loss is {P.outcome_logloss:.5f} versus {BASE.outcome_logloss:.5f}. These improvements concern regulation outcomes, not final overtime or shootout winners.',
  f'The marginal central intervals cover {P.home_coverage90:.1%} of home counts and {P.away_coverage90:.1%} of away counts, with mean width {P.mean_interval_width:.3f} goals. Coverage above the nominal level partly reflects discrete quantile endpoints. It should not be interpreted as evidence of sharp, perfectly calibrated uncertainty.',
  f'The paired bootstrap primary-minus-baseline NLL interval is [{PC.low95:.5f}, {PC.high95:.5f}], with mean {PC.mean_difference:.5f}, based on {int(PC.blocks)} calendar-week blocks. This supports an improvement within the observed holdout under the block scheme. It does not establish identical gains in another season or under unobserved future roster information.'],level=2)
 section(doc,'8 Calibration uncertainty and reproducibility',[
  'Calibration asks whether forecast probabilities agree with realized frequencies. Fixed ten-percent outcome bins are used here as descriptive diagnostics after evaluation. No calibration map is fitted to holdout outcomes. Sparse bins are retained with their counts in the accompanying CSV rather than treated as reliable estimates.',
  'The paired score model’s zero positive-dependence component is a limitation and a finding. Hockey score interaction can be negative, state dependent or poorly described by a shared count. Empty-net scoring and late-game strategy are mechanisms worth studying, but this dataset does not establish their causal contribution.',
  'The archive and raw-source fingerprints identify the exact local materials. Frozen model configurations, tuning choices, fold metrics, data summaries, exclusion records, score checks and comparison intervals trace the reported numbers to versioned artifacts. Row-level prediction files remain local, with aggregate findings published.',
  'The base-R independent-Poisson implementation uses the same features, training period, population-standard-deviation scaling, penalty and log-link likelihood as Python. Its final holdout joint NLL differs by approximately 1.71 billionths and its MAE by approximately 2.64 billionths. This is a targeted cross-language numerical check, not full R parity for every model.',
  'The frozen Python replay checks the exact source and protocol hashes before refitting existing configurations. It performs no hyperparameter search and does not revise the published holdout scores. Fast synthetic invariant tests can run in CI without access to restricted raw data. Reproducing the full empirical experiment also requires the original source snapshot; current downloads may change.'],page=True)
 figure(doc,'outcome-calibration.png','Figure 5. Descriptive reliability curves for the frozen model. Low-count bins are noisy; no holdout-based calibration adjustment was applied.')
 section(doc,'9 Discussion and limitations',[
  'The original thesis correctly identified that per-game modeling was needed. The revision operationalizes that insight and makes paired probability forecasts possible. It cannot quantify a clean causal effect of switching from team-season OLS to game-level counts because the historical dataset, original code and identical target definition are not available.',
  'The strongest evidence is that practical regularization and lagged game histories improve probability quality modestly under this temporal design. The league baseline remains competitive on point error. A small point-error reduction can coexist with a useful proper-score gain because a full distribution is evaluated differently from a point forecast.',
  'The correlated-score hypothesis is not supported by the shared-count model as fitted. Rho reaches zero, making the bivariate model operationally identical to the independent one. This outcome does not prove score independence; the model permits only one restricted form of positive dependence. A future comparison could examine negative-dependence or state-based alternatives under a new untouched validation period.',
  'There is no basis to claim that hierarchical Bayesian inference or dynamic state-space modeling has been completed. The implemented attack/defence model uses L2-penalized team indicators; it is a useful shrinkage benchmark but lacks posterior uncertainty and time-varying latent strengths. The tree models are controlled benchmarks, not a comprehensive machine-learning search.',
  'Historical availability is the most important deployment limitation. Strict date lags prevent direct target inclusion, but the present download is a revised modern snapshot. Exact provider release time, archived revisions, actual start times and announced lineups are absent. A live forecasting claim requires prospective snapshots and an explicit timestamp ledger.',
  'Other limitations include one excluded incomplete game, sample-based rather than exhaustive official score verification, calendar-day rather than exact-time rest, no playoff form, unmodeled neutral-site or international venue effects, cold starts for changed team labels, and no roster, starting-goaltender, injury or market information. These omissions affect the study’s information set and external validity.',
  'The pandemic-disrupted seasons remain distinct folds but are not separately modeled as causal rule changes. The single-season final holdout and 28 bootstrap week blocks limit comparison uncertainty. Multiple candidate comparisons are descriptive and their intervals are not multiplicity-adjusted. Independent editorial and academic review remain required before describing the entire programme as a completed formal thesis.'],page=True)
 section(doc,'10 Conclusion and remaining programme',[
  'This first research edition preserves the historical work, resolves key methodological definitions and completes a reproducible nine-variant temporal comparison. It produces valid paired count forecasts and reports a small, uncertainty-qualified probability-quality gain over an expanding league baseline.',
  'The negative findings are central: positive shared-count dependence is estimated at zero; larger tree models do not decisively improve the final comparison; and the gain in average absolute point error is very small. A revised thesis should retain these findings rather than substitute an unsupported narrative of algorithmic progress.',
  'The next scientific priorities are original artifact recovery, broader score verification, prospective historical-availability snapshots, feature exposure normalization, a carefully justified dynamic/hierarchical model, and development-only sensitivity work before any new test-season analysis. The reserved 2025-26 season must not become a tuning resource for the current release.',
  'The release is version 0.1.0, a substantive initial empirical revision. Version 1.0.0 requires the wider programme’s completion gates, including independent review and clearer archival status. No new institutional submission, portfolio website deployment, video recording or communication with Carleton is implied by this research release.'],page=True)
 section(doc,'References',REFERENCES,page=True)
 section(doc,'Appendix A Definitions and build trace',[
  'Team-game: one side’s goal count in one paired game. Game: one ordered home/away pair. Season labels use their starting year in machine-readable files. Regulation draw: equal goals after the first three periods, regardless of the final overtime or shootout result.',
  'Source-verified: an original reported value confirmed by inspecting its preserved source. Reproduced: obtained by executing preserved original code and original inputs. Reconstructed: recreated from documented methods with explicit replacement inputs or assumptions. These labels are not interchangeable.',
  f'Original archival PDF SHA256: {checksum(ROOT/"archive/Silvey-2020-NHL-Honours-Project.pdf")}',
  f'Modern raw snapshot SHA256: {F["source_sha256"]}',
  f'Frozen protocol SHA256: {F["protocol_sha256"]}',
  'Code entry points: scripts/acquire.py; scripts/validate_scores.py; scripts/run_study.py; scripts/reproduce_frozen.py; scripts/r_poisson_parity.R; scripts/build_figures.py; scripts/build_documents.py. Use the repository README for exact rebuild commands and current checks.',
  'Numerical source files: reports/data-summary.csv, development-metrics.csv, holdout-metrics.csv, holdout-paired-comparisons.csv, outcome-calibration.csv, frozen-models.json and r-parity.csv. All displayed table values are rounded only during document generation.'],page=True)
 section(doc,'Appendix B Frozen final configurations',[
  'The final hyperparameter choices were determined without final-season score comparison. They are retained verbatim in the frozen configuration ledger.'],level=1)
 table(doc,['Model','Frozen configuration'],[(NAMES[name],json.dumps(config,sort_keys=True)) for name,config in F['configs'].items()],[3.2,3.4])
 sensitivity=pd.read_csv(ROOT/'reports/development-sensitivity.csv')
 aggregate=sensitivity.groupby('variant').apply(lambda g:(g.joint_nll*g.n).sum()/g.n.sum(),include_groups=False)
 default=float(weighted['poisson'])
 section(doc,'Appendix C Exploratory development sensitivity',[
  'These follow-up checks were specified after the initial holdout evaluation. They access only season-start years 2016 through 2023 and do not change the frozen primary model, its configurations or any final-season score. Their status is exploratory rather than pre-specified confirmatory evidence.',
  'The same Poisson search penalties and inner last-training-season validation are used for recent-form windows of 5, 10 and 40 games, with the original 10-game priors unchanged. A separate 20-game variant removes historical shots and attempts while retaining goal form, rest, home role and league rate.',
  'A longer 40-game window has the lowest development NLL among these follow-up variants; the shorter windows are worse. Removing shot and attempt predictors increases development loss. This suggests that feature-history choices merit study, but the variants are not tested on the already-observed final season and are not promoted as a new primary model.',
  'Any future choice of window or feature subset must be frozen using development data and evaluated on a new untouched test period. The follow-up results do not justify reusing the 2024-25 outcomes for optimization.'],page=True)
 table(doc,['Development variant','Joint NLL','Difference from initial Poisson'],
       [('Initial 20-game full features',f'{default:.5f}','0.00000')]+
       [(name.replace('_',' '),f'{value:.5f}',f'{value-default:+.5f}') for name,value in aggregate.items()],[3.1,1.3,2.2])
 p(doc,'Table 6. Game-weighted exploratory development loss over the same five seasons and 5,885 games. The initial frozen holdout and model-selection record are unchanged.')
 save(doc,'NHL-Goal-Prediction-2026-Thesis-Revision')

def audit():
 doc=new_doc('NHL Honours Project Archival Audit','Source verified findings and reproduction status\nCompanion to the 2026 thesis revision')
 section(doc,'Audit conclusion',[
  'The 2020 honours project survives as a complete 47-page PDF. It is sufficient to verify the broad research design and principal reported numbers. Exact execution cannot yet be reproduced because the local recovery inventory has not found the original NHL data table or executable R analysis.',
  'The major methodological problem is the mismatch between an individual-game forecasting objective and a team-season estimation sample. The revision preserves the original record and corrects the unit, target, information boundary and validation design in a separately labeled modern experiment.',
  'This report distinguishes what the source says, what has been independently checked by inspection, what remains ambiguous and what must be recovered before an exact archival reproduction can be claimed.'])
 contents(doc,['1 Archival identity and preservation','2 Source recovery and evidence status','3 Numerical reconciliation',
  '4 Methodological corrections','5 Data dictionary corrections','6 Reproduction plan','7 Relation to the new experiment','References'])
 section(doc,'1 Archival identity and preservation',[
  'Author in the original title page: Howard Silvey. Supervisor: Dr. Jason Nielsen. Institution: Carleton University, School of Mathematics and Statistics. Submission date: April 7, 2020. Title: National Hockey League Data: Analysis and Prediction.',
  'The 2026 author identity is Howard R. Silvey. The new work is titled National Hockey League Goal Prediction, with the subtitle Modern Statistical and Machine Learning Approaches and a revision-and-extension statement. The 2020 submission date is not applied to the new analysis.',
  'The local archival PDF is copied without rewriting metadata, re-pagination or repairs. Its public Carleton URL is retained as provenance. The accidental suffix in the original public filename is not treated as a scientific version identifier.',
  f'Preserved artifact: archive/Silvey-2020-NHL-Honours-Project.pdf. Pages: 47. SHA256: {checksum(ROOT/"archive/Silvey-2020-NHL-Honours-Project.pdf")}.'],page=True)
 section(doc,'2 Source recovery and evidence status',[
  'The recovery scope for this edition is a targeted local filename inventory of University materials, Documents and Desktop Downloads. Matching NHL, hockey, honours, R source/history, spreadsheet and archive names were considered. No original NHL executable source or complete input table was recovered. The surviving PDF is an original artifact; the supplied 2026 baseline is a later derivative record.',
  'The Documents R history contains recent archive-extraction commands for a different project and does not supply NHL analysis. Unrelated private files were not copied into the research repository. This inventory is not a claim that every cloud drive, email attachment, removable device or older computer was searched.',
  'The PDF evidence permits source verification, not rerunning the analysis. Reproduction should remain unresolved until surviving original materials are found or their absence is explicitly established. A reconstruction is an acceptable scientific fallback only when its assumptions and discrepancies are disclosed.'],page=True)
 table(doc,['Artifact or finding','Classification','Current status'],[
  ('Original 2020 PDF','Original','Preserved and checksummed'),('September 2026 baseline','Derivative audit','Preserved with source attribution'),
  ('Original R scripts','Unknown','Not found in targeted local search'),('Original team-season table','Unknown','Not found in targeted local search'),
  ('Historical headline metrics','Reported original results','Visually source-verified, not recomputed'),
  ('Modern game dataset and models','New 2026 analysis','Separate empirical experiment')],[2.4,1.8,2.4])
 section(doc,'3 Numerical reconciliation',[
  'Physical PDF page numbers below count the cover as page 1. The document’s printed chapter pagination begins later and therefore differs. The main LOOCV screenshot is on physical page 39, printed page 32.',
  'The table dimensions support 123 team-season rows and approximately 40 candidate columns. The source’s 4,920 figure counts table cells, not the number of independent observations. The cross-validation display itself states 123 samples, five predictors and training samples of size 122.',
  'The backward model retains CF, CA, CF_Percent, PP_Percent and L. Its reported R squared is 0.1161. The forward model’s reported R squared is 0.06451. The source discusses adjusted R squared but the final console values and selection path still need exact code-level reconciliation; this audit does not infer a full coefficient table from narrative alone.'],page=True)
 table(doc,['Claim','Verified source','Execution status'],[
  ('123 estimation samples','Physical pages 17 and 39','Original rows unavailable'),('Backward R squared 0.1161','Physical pages 28 and 32','Not recomputed'),
  ('Forward R squared 0.06451','Physical page 32','Not recomputed'),('RMSE 46.20022','Physical page 39 screenshot','Not recomputed'),
  ('R squared 0.03927991','Physical page 39 screenshot','Not recomputed'),('MAE 37.91543','Physical page 39 screenshot','Not recomputed')],[2.2,2.,2.4])
 p(doc,'The LOOCV values concern the original season-total response. Dividing an RMSE by 82 would yield a rescaled season-average error, not a validated individual-game forecasting error. Modern regulation-goal metrics must not be presented as a direct before-and-after reduction on the same task.')
 section(doc,'4 Methodological corrections',[
  'Unit and target. Replace team-season totals with one paired record per game. Define the two goal counts explicitly and separate regulation, overtime and shootout conventions. The paired-response idea is already anticipated by the original abstract.',
  'Leakage. Exclude target-game and same-period outcome-derived features. Lag shooting, special-teams, save, standings and scoring measures. A valid feature can be unsafe when its time window overlaps the target period.',
  'Validation. Replace pooled random sampling and LOOCV with ordered expanding windows and tuning inside the training period. Team repetition and changing scoring environments require attention beyond a simple independence assumption.',
  'Outcome support. OLS can produce negative point estimates and supplies no count distribution without additional assumptions. Retain an historical OLS benchmark where reconstruction becomes possible; use count models for coherent probability forecasts.',
  'Selection and collinearity. CF, CA and CF percentage are algebraically related. Other totals and ratios may introduce rank deficiency or unstable estimates. NA coefficients require a rank and preprocessing audit. Stepwise p values are not protected from repeated-search uncertainty.',
  'Inference and uncertainty. State OLS assumptions precisely; do not use the Central Limit Theorem to justify all residual properties. Separate coefficient inference, point forecasting and distribution forecasting. Report uncertainty for model differences rather than a universal acceptable RMSE threshold.',
  'Schedule exposure. A fixed 82-game divisor is inappropriate for incomplete or disrupted seasons. Rates must use actual exposure, or the response should be defined directly at game level.',
  'Claims. Two correctly rounded forecasts in seven displayed cases are illustrative, not a stable general accuracy estimate. The revision evaluates every admitted held-out game and retains negative results.'],page=True)
 section(doc,'5 Data dictionary corrections',[
  'FA is labeled inconsistently in the original glossary and appears intended to represent Fenwick Against. The high-danger percentage label and formula need reconciliation. Save percentage must use saves rather than an ambiguous S already defined as team shots.',
  'PDO should distinguish its shooting and save components and specify its scale. Penalty-kill percentage requires an explicit opponent power-play scoring numerator and denominator. Totals, rates, percentages and adjusted statistics must be marked separately.',
  'The revised registry should include canonical name, formula, unit, side, aggregation level, source, window, lag rule, earliest availability, missing-value policy and transformation lineage. These definitions make a predictor auditable rather than merely recognizable to a hockey reader.'],page=True)
 table(doc,['Original issue','Required archival check','Revision treatment'],[
  ('FA label','Inspect source column and glossary','Canonical for/against naming'),('High-danger percentage','Recover exact inputs/formula','No unverified formula carried forward'),
  ('SV percentage ambiguity','Identify actual saves numerator','Explicit saves/shot exposure definition'),('PDO description','Verify scale and component sums','Separate definitions and availability'),
  ('Home goals versus team total','Recover response construction','Paired regulation target'),('Incomplete schedule','Recover extraction cutoff and games played','Game counts and exclusions explicit')],[1.8,2.4,2.4])
 section(doc,'6 Reproduction plan',[
  'First inventory surviving source files without modification. Calculate hashes, record original names and provenance, identify restricted materials and map every artifact to thesis tables and figures. Recovery can include a user-supplied archive without renaming its internal contents.',
  'Restore the R interpreter and packages as closely as evidence permits. Preserve the original analysis separately from any corrected re-analysis. If exact environment restoration is impossible, record version substitutions and assess their numerical impact.',
  'Verify the 123 rows, season labels, extraction cutoff and approximately 40 columns. Recover the true response, exposures, import types, missing values and deterministic feature identities. Reproduce the full OLS fit and inspect aliased coefficients before trying to match selection output.',
  'Recover both selection procedures, significance thresholds, exact formulas and coefficients. Reconstruct the three displayed prediction samples and determine whether the data used for each were genuinely excluded from estimation.',
  'Re-run caret LOOCV with selection and preprocessing placement documented. Reconcile each headline metric and every material discrepancy. A discrepancy ledger must state whether it comes from source transcription, a different sample, altered preprocessing, selection behavior or environment change.',
  'If original inputs cannot be recovered, build a clearly labeled reconstruction from equivalent historical sources. Never relabel reconstruction as reproduction. The absence of original files limits archival claims but does not invalidate a separately designed modern study.'],page=True)
 section(doc,'7 Relation to the new experiment',[
  'The modern study uses 10,969 paired regular-season records, reconstructed regulation goals and conservative earlier-date features. It implements nine variants with five development seasons and a separately frozen final season. Its results belong to the 2026 extension.',
  'The initial empirical result is a modest proper-score improvement from regularized recent-form predictors. The shared bivariate component is zero and the point-error gain is small. These results neither reproduce nor numerically correct the original LOOCV output.',
  'The archival status remains source-verified but not execution-reproduced. The project status ledger and release notes retain that qualification. This report is complete as an audit of available evidence, while the reproduction task remains open.'],page=True)
 section(doc,'References',[REFERENCES[0],REFERENCES[10],REFERENCES[11]])
 save(doc,'NHL-2020-Archival-Audit-and-Reproduction-Status')

def protocol():
 doc=new_doc('NHL Goal Prediction Research Protocol','Data contract temporal validation and reproducibility\nFrozen first empirical study with technical appendices')
 section(doc,'Protocol purpose',[
  'This document specifies the first empirical study independently of its observed performance. The repository protocol was committed before development comparisons. The appendix records the actual frozen model choices without modifying the earlier search spaces.',
  'The complete research programme is broader than this first release. This protocol defines what has been implemented and which advanced model families remain prospective. It supplies the conditions needed to interpret the thesis’s numerical comparison.'])
 text=(ROOT/'docs/PROTOCOL.md').read_text(encoding='utf-8')
 paragraphs=[]
 current=[]
 for line in text.splitlines():
  if line.startswith('# '):continue
  if line.startswith('## '):
   if current:p(doc,' '.join(current));current=[]
   doc.add_heading(line[3:],1)
  elif re.match(r'^\d+\.',line):
   if current:p(doc,' '.join(current));current=[]
   current.append(line.strip())
  elif not line.strip():
   if current:p(doc,' '.join(current));current=[]
  else:current.append(line.strip())
 if current:p(doc,' '.join(current))
 section(doc,'Appendix A Machine readable data dictionary',[
  'The canonical table preserves source game ID, season start year, source date, home and away team labels, regulation goals, actual played goals, duration, overtime/shootout indicators, and prior-game shot/attempt totals. The analytical feature table adds strictly earlier-date histories.',
  'date is a source calendar-date boundary, not a verified UTC publication timestamp. history_date is the latest source date admitted for the team. The source-date invariant requires history_date < date. A missing first history is represented as no prior observation, with a fixed pooled prior.'],page=True)
 table(doc,['Field or family','Definition','Availability or policy'],[
  ('game_id and season','Stable source game ID; starting-year season','One unique paired row'),('home_goals and away_goals','Reconstructed regulation counts','Outcome only; never a current predictor'),
  ('duration','All-situation elapsed seconds','Target reconstruction only'),('played_goals','Actual goals including overtime','Exclude shootout scoreboard marker'),
  ('team_gf and team_ga','Smoothed prior 20-game regulation goals','Earlier source dates only'),('team_sf and team_sa','Smoothed prior-game shots','Overtime exposure may remain'),
  ('team_cf and opp_ca','Smoothed prior-game shot attempts','Explicit lag; correlated predictors regularized'),('rest and opp_rest','Days since prior included game; cap 7','Calendar-date proxy'),
  ('league_rate','Expanding side-specific earlier-date mean','100 fixed prior games'),('team and opponent','Categorical team labels','Unseen category pooled in model')],[1.9,2.4,2.3])
 section(doc,'Appendix B Model and information cards',[
  'Model scope: historical NHL regular-season paired regulation counts. Intended use: non-commercial comparative research and education. Excluded use claims: guaranteed exact scores, causal intervention effects, betting profitability, final winner probabilities without an overtime model, or institutional acceptance.',
  'Primary configuration: the frozen bivariate-Poisson label, with estimated rho zero, produces the independent-Poisson distribution. Independent and bivariate variants are numerically tied. A future tie-selection rule should explicitly prefer the simpler model, but no post-holdout relabeling changes this release’s frozen record.',
  'Information set: historical goals, shots, attempts, source-date rest and home role. Roster, injuries, confirmed goalies, odds, xG, player tracking, travel and exact venue context are not part of the primary model. Their absence must not be interpreted as proof that they lack value.',
  'Known limitations: revised retrospective source snapshots; calendar-time proxy; sample-based score verification; one excluded incomplete game; limited final-season horizon; positive-only shared-count dependence; discrete interval overcoverage; proxy distributions for squared-error estimators.'])
 section(doc,'Appendix C Frozen configurations and environment',[],page=True)
 table(doc,['Model','Frozen final choice'],[(NAMES[n],json.dumps(c,sort_keys=True)) for n,c in F['configs'].items()],[3.1,3.5])
 p(doc,f'Python interpreter used for fitting: {F["python"]}. Random seed: {F["seed"]}. Training is deterministic under the stated software and numerical tolerance, with four computation threads used for reproducible bounded parallelism.')
 table(doc,['Package','Recorded version'],[(n,v) for n,v in F['packages'].items()],[3.1,3.5])
 p(doc,'R parity uses base R 4.2.3 and no contributed packages. Its interpreter and session record define that environment. A broader contributed-package R suite would require its own dependency lock. Python dependency pins and a generated full dependency lock document the actual environment.')
 section(doc,'Appendix D Verification and release gates',[
  'Fast verification checks count support, score reconstruction edge cases, future-outcome mutation invariance, same-date batching, side stacking, bivariate mass/moments, outcome mass and bootstrap pairing. CI can execute these on synthetic data without network access or third-party raw downloads.',
  'Full replay validates the exact raw and protocol hashes, refits only frozen configurations and requires aggregate holdout metrics to agree within 1e-7. It does not select hyperparameters again. The R Poisson parity threshold is 1e-6.',
  'Publication retains the original PDF, provenance hashes, aggregate tables, model choices, protocol, tests, manuscript sources and final documents. Raw and row-level outputs remain excluded. Software licensing does not confer rights over MoneyPuck or NHL content.',
  'A stable 1.0.0 research release still requires resolved archival status, independent review, broader score and availability verification, sensitivity analysis and any justified advanced models. The current 0.1.0 release must keep that remaining work visible.'])
 section(doc,'References',[REFERENCES[1],REFERENCES[2],REFERENCES[7],REFERENCES[8],REFERENCES[9],REFERENCES[11]])
 save(doc,'NHL-2026-Research-Protocol-and-Reproducibility')

if __name__=='__main__':
 thesis();audit();protocol()
