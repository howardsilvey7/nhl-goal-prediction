"""One-command public release integrity check without third-party raw data."""
from pathlib import Path
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))
import json
import subprocess
import zipfile
import numpy as np
import pandas as pd
from pypdf import PdfReader
from docx import Document
from nhl_forecast.data import checksum

def main():
    subprocess.run([sys.executable,'-m','unittest','discover','-s',str(ROOT/'tests'),'-v'],check=True)
    frozen=json.loads((ROOT/'reports/frozen-models.json').read_text())
    if checksum(ROOT/'docs/PROTOCOL.md')!=frozen['protocol_sha256']:
        raise ValueError('Frozen protocol hash mismatch')
    archive=json.loads((ROOT/'archive/manifest.json').read_text())
    if checksum(ROOT/'archive'/archive['filename'])!=archive['sha256']:
        raise ValueError('Original PDF changed')
    data=pd.read_csv(ROOT/'reports/data-summary.csv')
    if data.n_games.sum()!=10969:raise ValueError('Canonical count mismatch')
    hold=pd.read_csv(ROOT/'reports/holdout-metrics.csv')
    dev=pd.read_csv(ROOT/'reports/development-metrics.csv')
    if len(hold)!=9 or not (hold.n==1312).all():raise ValueError('Holdout dimensions mismatch')
    if len(dev)!=45 or set(dev.season)!={2019,2020,2021,2022,2023}:raise ValueError('Temporal fold mismatch')
    if not np.isfinite(hold.select_dtypes('number').to_numpy()).all():raise ValueError('Invalid numeric result')
    score=pd.read_csv(ROOT/'reports/score-validation.csv')
    if len(score)!=54 or not score.passed.all():raise ValueError('Official score verification failed')
    manifests=json.loads((ROOT/'reports/document-manifest.json').read_text())
    for entry in manifests:
        path=ROOT/'outputs'/entry['filename']
        if checksum(path)!=entry['sha256']:raise ValueError('Document checksum mismatch')
        if path.suffix=='.pdf':
            pdf=PdfReader(path)
            if len(pdf.pages)!=entry['pages']:raise ValueError('PDF pagination mismatch')
            text='\n'.join(p.extract_text() or '' for p in pdf.pages)
            for token in ['TODO','[INSERT','turn0search','C:\\Users\\','Error!']:
                if token in text:raise ValueError(f'Unfinished or private document token: {token}')
        else:
            doc=Document(path)
            if not doc.paragraphs or doc.paragraphs[0].style.name!='Title':raise ValueError('Missing title semantics')
            if any(True for _ in doc.styles['Title']._element.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr')):
                raise ValueError('Unwanted title border')
    print('Release integrity verified: archive, protocol, 45 development rows, nine holdout variants, 54 score checks and six final documents.')

if __name__=='__main__':main()
